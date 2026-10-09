"""扫描 DOCX，确认没有残留的 Markdown / LaTeX 痕迹。"""
import re
import sys
from docx import Document

DOCX = sys.argv[1] if len(sys.argv) > 1 else r"D:\myagent\选题说明.docx"

PATTERNS = {
    "美元符号 $": re.compile(r"\$"),
    "粗体标记 **": re.compile(r"\*\*"),
    "链接语法 ](": re.compile(r"\]\("),
    "LaTeX \\ge / \\le": re.compile(r"\\[lg]e\b"),
    "LaTeX \\binom": re.compile(r"\\binom"),
    "反引号 `": re.compile(re.escape("`")),
    "LaTeX 其余反斜杠命令": re.compile(r"\\[a-zA-Z]{2,}"),
}

doc = Document(DOCX)
texts = [p.text for p in doc.paragraphs]
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            texts.append(cell.text)

hits = {name: [] for name in PATTERNS}
for text in texts:
    for name, pat in PATTERNS.items():
        if pat.search(text):
            hits[name].append(text)

print(f"段落/单元格总数: {len(texts)}")
total = 0
for name, items in hits.items():
    total += len(items)
    flag = "OK " if not items else "!! "
    print(f"{flag}{name}: {len(items)}")
    for text in items[:5]:
        print(f"      > {text[:110]}")

print("残留问题总数:", total)
sys.exit(0)
